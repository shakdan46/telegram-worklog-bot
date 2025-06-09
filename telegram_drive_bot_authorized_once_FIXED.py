
# telegram_drive_bot_authorized_once.py
import os
import io
import json
from datetime import datetime, timedelta
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (ApplicationBuilder, CommandHandler, CallbackQueryHandler,
                          ContextTypes, ConversationHandler, MessageHandler, filters)
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from openpyxl import load_workbook
import tempfile
from dotenv import load_dotenv

load_dotenv()

# שלבים
SELECTING_DATE, SELECTING_WORKERS, CONFIRM_SELECTION, ASK_NAME, ASK_SALARY, ASK_START_DATE = range(6)

# הגדרות
SCOPES = ['https://www.googleapis.com/auth/drive']
FILE_ID = os.environ.get("EXCEL_FILE_ID")
AUTHORIZED_USERS_FILE = "authorized_users.json"
ADMIN_PASSWORD = "204560916"
CREDENTIALS_JSON = json.loads(os.environ.get("GOOGLE_APPLICATION_CREDENTIALS_JSON"))

credentials = service_account.Credentials.from_service_account_info(CREDENTIALS_JSON, scopes=SCOPES)
drive_service = build('drive', 'v3', credentials=credentials)

MONTH_MAP = {
    1: "ינואר", 2: "פברואר", 3: "מרץ", 4: "אפריל", 5: "מאי", 6: "יוני",
    7: "יולי", 8: "אוגוסט", 9: "ספטמבר", 10: "אוקטובר", 11: "נובמבר", 12: "דצמבר"
}

def download_excel():
    request = drive_service.files().get_media(fileId=FILE_ID)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        status, done = downloader.next_chunk()
    fh.seek(0)
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    temp_file.write(fh.read())
    temp_file.close()
    return temp_file.name

def upload_excel(file_path):
    media_body = MediaIoBaseUpload(open(file_path, 'rb'), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    drive_service.files().update(fileId=FILE_ID, media_body=media_body).execute()

def is_authorized(user_id):
    if os.path.exists(AUTHORIZED_USERS_FILE):
        with open(AUTHORIZED_USERS_FILE, "r") as f:
            return str(user_id) in f.read().splitlines()
    return False

def authorize_user(user_id):
    with open(AUTHORIZED_USERS_FILE, "a") as f:
        f.write(f"{user_id}\n")

# התחלה ואימות סיסמה
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    if is_authorized(user_id):
        await update.message.reply_text("שלום! אנא הזן תאריך בפורמט YYYY-MM-DD")
        return SELECTING_DATE
    else:
        await update.message.reply_text("🔐 אנא הזן סיסמה:")
        return ASK_NAME

async def ask_password(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text.strip() == ADMIN_PASSWORD:
        authorize_user(update.effective_user.id)
        await update.message.reply_text("✅ הצלחה! אנא הזן תאריך בפורמט YYYY-MM-DD")
        return SELECTING_DATE
    else:
        await update.message.reply_text("❌ סיסמה שגויה.")
        return ConversationHandler.END

async def ask_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["worker_name"] = update.message.text.strip()
    await update.message.reply_text("מה השכר היומי?")
    return ASK_SALARY

async def ask_salary(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        context.user_data["salary"] = float(update.message.text.strip())
        await update.message.reply_text("תאריך התחלה (dd/mm/yyyy):")
        return ASK_START_DATE
    except:
        await update.message.reply_text("סכום לא תקין. נסה שוב:")
        return ASK_SALARY

async def ask_start_date(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        start_date = datetime.strptime(update.message.text.strip(), '%d/%m/%Y')
    except:
        await update.message.reply_text("תאריך שגוי. נסה שוב:")
        return ASK_START_DATE

    name = context.user_data["worker_name"]
    salary = context.user_data["salary"]

    file_path = download_excel()
    wb = load_workbook(file_path)

    current_date = start_date
    while current_date.year == 2025:
        sheet_name = MONTH_MAP[current_date.month]
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            row = sheet.max_row + 1
            sheet.cell(row=row, column=1).value = current_date.strftime('%d/%m/%Y')
            sheet.cell(row=row, column=2).value = name
            sheet.cell(row=row, column=5).value = salary
        current_date += timedelta(days=1)

    if "יומית פועלים" in wb.sheetnames:
        ysheet = wb["יומית פועלים"]
        ysheet.append([name, salary])

    wb.save(file_path)
    upload_excel(file_path)
    os.remove(file_path)

    await update.message.reply_text("✅ פועל נוסף בהצלחה.")
    return ConversationHandler.END

def main():
    TOKEN = os.environ["TELEGRAM_TOKEN"]
    app = ApplicationBuilder().token(TOKEN).build()

    conv_add_worker = ConversationHandler(
        entry_points=[CommandHandler("addworker", start)],
        states={
            ASK_NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_name)],
            ASK_SALARY: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_salary)],
            ASK_START_DATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_start_date)],
        },
        fallbacks=[],
    )

    app.add_handler(conv_add_worker)
    app.run_polling()

if __name__ == "__main__":
    main()

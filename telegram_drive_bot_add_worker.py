import os
import io
import tempfile
from datetime import datetime, timedelta
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ApplicationBuilder, CommandHandler, CallbackQueryHandler,
    ContextTypes, ConversationHandler, MessageHandler, filters
)
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaFileUpload
from openpyxl import load_workbook

# שלבים
SELECTING_DATE, SELECTING_WORKERS, VERIFY_SELECTION, ASK_NAME, ASK_SALARY, ASK_DATE = range(6)

# הגדרות
SCOPES = ['https://www.googleapis.com/auth/drive']
SERVICE_ACCOUNT_FILE = 'credentials.json'  # שנה לקובץ הסרוויס שלך
SPREADSHEET_FILE_ID = '1UT4zWtuny8ES2z3dF2UTudK86HF1Mm5B'  # שנה ל-ID של קובץ הדרייב שלך

MONTHS = ['ינואר', 'פברואר', 'מרץ', 'אפריל', 'מאי', 'יוני',
          'יולי', 'אוגוסט', 'ספטמבר', 'אוקטובר', 'נובמבר', 'דצמבר']

def download_excel():
    credentials = service_account.Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=SCOPES)
    service = build('drive', 'v3', credentials=credentials)
    request = service.files().get_media(fileId=SPREADSHEET_FILE_ID)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        status, done = downloader.next_chunk()
    fh.seek(0)
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    temp_file.write(fh.read())
    temp_file.close()
    return temp_file.name

def upload_excel(file_path):
    credentials = service_account.Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=SCOPES)
    service = build('drive', 'v3', credentials=credentials)
    media_body = MediaFileUpload(file_path, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    service.files().update(fileId=SPREADSHEET_FILE_ID, media_body=media_body).execute()

# התחלת שיחה של הוספת פועל
async def start_add_worker(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("מה שם הפועל החדש?")
    return ASK_NAME

async def ask_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["worker_name"] = update.message.text.strip()
    await update.message.reply_text("מה השכר היומי שלו?")
    return ASK_SALARY

async def ask_salary(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        salary = float(update.message.text.strip())
    except ValueError:
        await update.message.reply_text("הזן סכום תקין (מספר בלבד).")
        return ASK_SALARY
    context.user_data["salary"] = salary
    await update.message.reply_text("מה תאריך התחלת העבודה של הפועל? (dd/mm/yyyy)")
    return ASK_DATE

async def ask_date(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        start_date = datetime.strptime(update.message.text.strip(), '%d/%m/%Y')
    except ValueError:
        await update.message.reply_text("הזן תאריך בפורמט dd/mm/yyyy (לדוגמה: 15/06/2025)")
        return ASK_DATE

    name = context.user_data["worker_name"]
    salary = context.user_data["salary"]
    file_path = download_excel()
    wb = load_workbook(file_path)

    month_name_map = {
        1: "ינואר", 2: "פברואר", 3: "מרץ", 4: "אפריל", 5: "מאי", 6: "יוני",
        7: "יולי", 8: "אוגוסט", 9: "ספטמבר", 10: "אוקטובר", 11: "נובמבר", 12: "דצמבר"
    }

    current_date = start_date
    while current_date.year == 2025:
        sheet_name = month_name_map[current_date.month]
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            row = sheet.max_row + 1
            sheet.cell(row=row, column=1).value = current_date.strftime('%d/%m/%Y')
            sheet.cell(row=row, column=2).value = name
            sheet.cell(row=row, column=5).value = salary
        current_date += timedelta(days=1)

    # עדכון גם ביומית פועלים
    if "יומית פועלים" in wb.sheetnames:
        sheet = wb["יומית פועלים"]
        row = sheet.max_row + 1
        sheet.cell(row=row, column=1).value = name
        sheet.cell(row=row, column=2).value = salary

    wb.save(file_path)
    upload_excel(file_path)
    os.remove(file_path)

    await update.message.reply_text(f"✅ הפועל {name} נוסף מהתאריך {start_date.strftime('%d/%m/%Y')} עם שכר יומי של ₪{salary}.")
    return ConversationHandler.END

# ביטול
async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("הפעולה בוטלה.")
    return ConversationHandler.END

def main():
    app = ApplicationBuilder().token("YOUR_BOT_TOKEN").build()

    conv_handler_add_worker = ConversationHandler(
        entry_points=[CommandHandler("addworker", start_add_worker)],
        states={
            ASK_NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_name)],
            ASK_SALARY: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_salary)],
            ASK_DATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_date)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
    )

    app.add_handler(conv_handler_add_worker)
    app.run_polling()

if __name__ == '__main__':
    main()
